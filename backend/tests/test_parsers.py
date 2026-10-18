from io import BytesIO

from docx import Document
from pypdf import PdfWriter

from job_evidence_agent.parsers.docx import parse_docx
from job_evidence_agent.parsers.markdown import parse_markdown
from job_evidence_agent.parsers.pdf import PdfExtractionError, parse_pdf
from job_evidence_agent.parsers.text import parse_text


def test_markdown_and_text_keep_human_readable_locators() -> None:
    markdown_chunks = parse_markdown("# 技术栈\nFastAPI\n## 部署\nDocker Compose")
    text_chunks = parse_text("普通文本材料")

    assert [(chunk.text, chunk.locator) for chunk in markdown_chunks] == [
        ("FastAPI", "技术栈"),
        ("Docker Compose", "部署"),
    ]
    assert [(chunk.text, chunk.locator) for chunk in text_chunks] == [
        ("普通文本材料", "全文")
    ]


def test_docx_generates_paragraph_locators() -> None:
    document = Document()
    document.add_heading("项目经历", level=1)
    document.add_paragraph("使用 FastAPI 实现服务。")
    output = BytesIO()
    document.save(output)

    chunks = parse_docx(output.getvalue())

    assert [(chunk.text, chunk.locator) for chunk in chunks] == [
        ("项目经历", "第 1 段"),
        ("使用 FastAPI 实现服务。", "第 2 段"),
    ]


def test_pdf_with_no_extractable_text_returns_controlled_error() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = BytesIO()
    writer.write(output)

    try:
        parse_pdf(output.getvalue())
    except PdfExtractionError as error:
        assert str(error) == "该 PDF 没有可提取的文本，请改用可复制文本 PDF 或粘贴内容"
    else:
        raise AssertionError("不可提取文本的 PDF 应被拒绝")


def test_copyable_text_pdf_generates_page_locator() -> None:
    chunks = parse_pdf(make_text_pdf("Hello PDF"))

    assert [(chunk.text, chunk.locator) for chunk in chunks] == [("Hello PDF", "第 1 页")]


def make_text_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, content in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode())
        output.extend(content)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    output.extend(b"".join(f"{offset:010} 00000 n \n".encode() for offset in offsets[1:]))
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode()
    )
    return bytes(output)
